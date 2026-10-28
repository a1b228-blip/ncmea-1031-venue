"""產生「林口長庚場地 × 考場配置圖」結合示意圖（linkou.html）。

讀 NCMEA 參賽說明第 2 頁的平面圖疊加標籤（含紅色封鎖線），換算到原始平面圖像素座標，
再疊上「考場配置圖（第 3 頁）」各區塊對應到模擬教室(一)(二)的位置。
用法：python3 generate_linkou_mapping.py <NCMEA醫事放射組第二場.pptx>
"""
import html
import math
import sys
from pathlib import Path

from pptx import Presentation

OUT = Path(__file__).with_name('linkou.html')
E = 12700
IMG_W, IMG_H = 960, 620                  # 林口長庚平面圖.jpg 原始像素
PIC = (137, 120, 823, 500)               # 平面圖在投影片上的位置（pt，已含群組換算）
NS = '{http://schemas.openxmlformats.org/drawingml/2006/main}'

# 投影片顏色（主題色以 Office 預設值替代）
THEME = {'accent1': '#4472C4', 'accent2': '#ED7D31', 'accent3': '#A5A5A5', 'bg1': '#FFFFFF', 'lt1': '#FFFFFF'}


def to_px(x, y):
    return ((x - PIC[0]) / (PIC[2] - PIC[0]) * IMG_W, (y - PIC[1]) / (PIC[3] - PIC[1]) * IMG_H)


def group_xfrm(g):
    x = g._element.grpSpPr.find(NS + 'xfrm')
    o, e, co, ce = (x.find(NS + t) for t in ('off', 'ext', 'chOff', 'chExt'))
    return int(o.get('x')), int(o.get('y')), int(e.get('cx')), int(e.get('cy')), \
        int(co.get('x')), int(co.get('y')), int(ce.get('cx')), int(ce.get('cy'))


def fill_of(sh):
    sp = sh._element.spPr
    sf = sp.find(NS + 'solidFill')
    if sf is None:
        return None
    c = sf[0]
    v = c.get('val')
    return '#' + v if c.tag.endswith('srgbClr') else THEME.get(v, '#cccccc')


def collect(slide):
    """回傳第 2 頁所有疊加圖形：(名稱, 文字, 填色, 四角像素座標)。"""
    out = []

    def walk(shapes, T):
        for sh in shapes:
            if sh.shape_type == 6:
                ox, oy, ex, ey, cx, cy, cex, cey = group_xfrm(sh)
                sx, sy = ex / cex, ey / cey
                walk(sh.shapes, lambda x, y, T=T, ox=ox, oy=oy, cx=cx, cy=cy, sx=sx, sy=sy:
                     T(ox + (x - cx) * sx, oy + (y - cy) * sy))
                continue
            if sh.shape_type == 13:
                continue
            x1, y1 = T(sh.left, sh.top)
            x2, y2 = T(sh.left + sh.width, sh.top + sh.height)
            x1, y1, x2, y2 = x1 / E, y1 / E, x2 / E, y2 / E
            if not (PIC[0] - 5 < (x1 + x2) / 2 < PIC[2] + 5 and PIC[1] - 5 < (y1 + y2) / 2 < PIC[3] + 5):
                continue                                   # 標題、頁碼不畫
            cxm, cym, a = (x1 + x2) / 2, (y1 + y2) / 2, math.radians(sh.rotation)
            corners = []
            for px, py in ((x1, y1), (x2, y1), (x2, y2), (x1, y2)):  # 先在投影片座標旋轉再換算像素
                dx, dy = px - cxm, py - cym
                corners.append(to_px(cxm + dx * math.cos(a) - dy * math.sin(a),
                                     cym + dx * math.sin(a) + dy * math.cos(a)))
            text = sh.text_frame.text.replace('\x0b', '\n') if sh.has_text_frame else ''
            out.append((sh.name, text, fill_of(sh), corners))

    walk(slide.shapes, lambda x, y: (x, y))
    return out


def poly(c, **attrs):
    pts = ' '.join(f'{x:.1f},{y:.1f}' for x, y in c)
    a = ' '.join(f'{k.replace("_", "-")}="{v}"' for k, v in attrs.items())
    return f'<polygon points="{pts}" {a}/>'


def text_svg(x, y, s, size, color='#1f2a37', weight=700, anchor='middle'):
    lines = s.split('\n')
    spans = ''.join(f'<tspan x="{x:.1f}" dy="{0 if i == 0 else size * 1.15:.1f}">{html.escape(l)}</tspan>'
                    for i, l in enumerate(lines))
    y0 = y - (len(lines) - 1) * size * 1.15 / 2
    return (f'<text x="{x:.1f}" y="{y0:.1f}" font-size="{size}" font-weight="{weight}" fill="{color}" '
            f'text-anchor="{anchor}" dominant-baseline="central">{spans}</text>')


def slide_overlay(shapes):
    """依投影片原樣重畫疊加標籤；紅色封鎖線另外加粗並標註。"""
    g, blocks = [], []
    for name, text, fill, c in shapes:
        xs, ys = [p[0] for p in c], [p[1] for p in c]
        w, h = max(xs) - min(xs), max(ys) - min(ys)
        if name.startswith('減號'):
            blocks.append(c)
            continue
        if text in ZONE_NAMES:                         # 這三區改由 zone_layer 標示，避免文字重疊
            continue
        if name.startswith('箭號'):
            g.append(poly(c, fill='#e53935'))
            continue
        if name.startswith('文字方塊'):
            col = '#e53935'
            bg = '#ffff00' if '入闈' in text else None
            if bg:
                g.append(poly(c, fill=bg, opacity=.9))
            g.append(text_svg(sum(xs) / 4, sum(ys) / 4, text, 11, col))
            continue
        g.append(poly(c, fill=fill or 'none', stroke='#4472C4', stroke_width=.8))
        if text:
            size = min(10.5, h * .62 / max(1, text.count('\n') + 1))
            g.append(text_svg(sum(xs) / 4, sum(ys) / 4, text, round(size, 1), weight=600))
    return g, blocks


# 考場配置圖（第 3 頁）各區 → 林口平面圖像素位置（由原始平面圖量得）
ZONES = {
    'A':    (294, 205, 437, 272),   # 模擬教室(一)＝原開刀房模擬教室
    'hall': (294, 273, 327, 303),   # 中央影音控制室西端的小通道
    'ctrl': (327, 273, 437, 303),   # 中央影音控制室（原 ER 控制室）
    'wait': (294, 303, 343, 372),   # 模擬教室(二) 西側
    'B':    (343, 303, 437, 372),   # 模擬教室(二) 其餘
}
DOOR = (292, 247)                    # 模擬教室(一) 西牆南段：依考場配置圖直通檢查室 A
ZONE_NAMES = {'模擬教室(一)', '模擬教室(二)', '中央影音控制室'}


def zone_layer():
    z, out = ZONES, []
    def rect(k, fill, stroke, dash=''):
        x1, y1, x2, y2 = z[k]
        d = f' stroke-dasharray="{dash}"' if dash else ''
        out.append(f'<rect x="{x1}" y="{y1}" width="{x2 - x1}" height="{y2 - y1}" fill="{fill}" '
                   f'stroke="{stroke}" stroke-width="2"{d}/>')
    rect('A', 'rgba(199,215,247,.95)', '#1d4ed8')
    rect('B', 'rgba(248,203,173,.95)', '#c2410c')
    rect('wait', 'rgba(197,224,180,.95)', '#15803d')
    rect('hall', 'rgba(221,198,245,.97)', '#7e22ce')
    rect('ctrl', 'url(#hatch)', '#475569', '4 3')
    c = lambda k: ((z[k][0] + z[k][2]) / 2, (z[k][1] + z[k][3]) / 2)
    out.append(text_svg(*c('A'), '檢查室 A\n模擬教室(一)', 12, '#1e3a8a', 800))
    out.append(text_svg(*c('B'), '檢查室 B\n模擬教室(二)', 11, '#7c2d12', 800))
    out.append(text_svg(*c('wait'), '等候區', 10, '#14532d', 800))
    out.append(text_svg(c('hall')[0], c('hall')[1], '走\n道', 7.5, '#581c87', 800))
    out.append(text_svg(*c('ctrl'), '中央影音控制室', 8.5, '#334155', 700))
    out.append('<rect x="287" y="318" width="4" height="27" fill="#e53935"/>')     # 等候區外門：封閉
    out.append(text_svg(284, 312, '外門封閉', 7, '#b91c1c', 800, 'end'))
    x, y = DOOR
    out.append(f'<path d="M{x - 28} {y} L{x - 3} {y}" stroke="#e53935" stroke-width="4" marker-end="url(#ah)"/>')
    out.append(text_svg(x - 30, y - 9, '進出大門', 8, '#b91c1c', 800, 'end'))
    return out


def block_layer(blocks):
    out = []
    for i, c in enumerate(blocks, 1):
        out.append(poly(c, fill='#e53935', stroke='#7f1d1d', stroke_width=1))
        xs, ys = [p[0] for p in c], [p[1] for p in c]
        out.append(f'<circle cx="{sum(xs) / 4:.1f}" cy="{sum(ys) / 4:.1f}" r="13" fill="none" '
                   f'stroke="#e53935" stroke-width="2" stroke-dasharray="3 2"/>')
        out.append(text_svg(sum(xs) / 4 + 15, sum(ys) / 4 - 13, f'封鎖{i}', 9, '#b91c1c', 800, 'start'))
    return out


def build(pptx):
    shapes = collect(Presentation(pptx).slides[1])
    labels, blocks = slide_overlay(shapes)
    defs = '''<defs>
  <pattern id="hatch" width="6" height="6" patternUnits="userSpaceOnUse" patternTransform="rotate(45)">
    <rect width="6" height="6" fill="#e2e8f0"/><line x1="0" y1="0" x2="0" y2="6" stroke="#64748b" stroke-width="1.5"/>
  </pattern>
  <marker id="ah" viewBox="0 0 10 10" refX="5" refY="5" markerWidth="4" markerHeight="4" orient="auto-start-reverse">
    <path d="M0 0 L10 5 L0 10 z" fill="#e53935"/></marker>
</defs>'''
    base = (f'<image href="林口長庚平面圖.jpg" width="{IMG_W}" height="{IMG_H}"/>'
            f'<g>{"".join(labels)}</g>')
    full = (f'<svg viewBox="0 0 {IMG_W} {IMG_H}" class="plan">{defs}{base}'
            f'<rect x="286" y="176" width="160" height="204" fill="none" stroke="#1d4ed8" stroke-width="3" '
            f'stroke-dasharray="8 5" rx="6"/>'
            f'{"".join(zone_layer())}{"".join(block_layer(blocks))}</svg>')
    zoom = (f'<svg viewBox="250 140 220 250" class="zoom">{defs}{base}'
            f'{"".join(zone_layer())}{"".join(block_layer(blocks))}</svg>')
    page = TEMPLATE.replace('{{FULL}}', full).replace('{{ZOOM}}', zoom)
    OUT.write_text(page, encoding='utf-8')
    print('寫入', OUT, '封鎖線', len(blocks), '條')


TEMPLATE = Path(__file__).with_name('linkou_template.html').read_text(encoding='utf-8')

if __name__ == '__main__':
    build(sys.argv[1])
